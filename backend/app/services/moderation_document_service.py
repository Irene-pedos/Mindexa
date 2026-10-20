"""
app/services/moderation_document_service.py

Phase 3: Document generation for moderation forms, exam papers, and marking schemes.
Uses python-docx and pure-Python qrcode.
Documents generated:
    1. Moderation Form (assessment, course, round, checklist results, comments, decision, moderator, date)
    2. Exam Paper (header, instructions, sections A/B/C with choice notes, questions, seal block with QR code)
    3. Marking Scheme (assessment, course, question answer keys and marking guidance)
"""

from __future__ import annotations

import io
from datetime import datetime, UTC
from typing import Any, Optional

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

try:
    import qrcode
except ImportError:  # pragma: no cover
    qrcode = None

from app.db.enums import QuestionType
from app.services.assessment_marks import _effective_aq_marks


class ModerationDocumentService:
    """Generates institutional DOCX documents for assessments and moderation."""

    @classmethod
    def generate_moderation_form(
        cls,
        assessment: Any,
        review: Any,
        moderator_name: str = "Moderator",
    ) -> bytes:
        """
        Generate DOCX Moderation Form.
        Includes: assessment details, course, round, checklist results, comments, decision note, signature block.
        """
        doc = Document()

        # Title
        p_title = doc.add_heading(level=0)
        run_title = p_title.add_run("INSTITUTIONAL MODERATION FORM")
        run_title.font.size = Pt(18)
        run_title.bold = True
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Subtitle
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_sub = p_sub.add_run("MINDEXA ACADEMIC ASSESSMENT QUALITY ASSURANCE")
        run_sub.font.size = Pt(10)
        run_sub.font.italic = True

        doc.add_paragraph()

        # Metadata table
        table = doc.add_table(rows=5, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"

        course_code = getattr(assessment, "course_code", "N/A") or "N/A"
        course_name = getattr(assessment, "course_name", "N/A") or "N/A"
        academic_year = getattr(assessment, "academic_year", "N/A") or "N/A"
        round_no = getattr(review, "round_no", 1) if review else 1
        review_status = (
            getattr(review.status, "value", str(review.status))
            if review and hasattr(review, "status")
            else "OPEN"
        )

        rows = [
            ("Assessment Title:", getattr(assessment, "title", "Untitled")),
            ("Course:", f"{course_code} - {course_name}"),
            ("Academic Year:", str(academic_year)),
            ("Moderation Round:", f"Round {round_no}"),
            ("Review Status:", str(review_status)),
        ]

        for i, (label, val) in enumerate(rows):
            cell_lbl = table.cell(i, 0)
            cell_val = table.cell(i, 1)
            cell_lbl.text = label
            cell_lbl.paragraphs[0].runs[0].bold = True
            cell_val.text = str(val)

        doc.add_paragraph()

        # Checklist Results Section
        doc.add_heading("Automated Checklist Results", level=2)
        checklist_res = getattr(review, "checklist_result", None) or {}
        items = checklist_res.get("items", []) if isinstance(checklist_res, dict) else []

        if items:
            chk_table = doc.add_table(rows=len(items) + 1, cols=4)
            chk_table.style = "Table Grid"
            headers = ["Check Code", "Severity", "Result", "Details"]
            for col_idx, h in enumerate(headers):
                cell = chk_table.cell(0, col_idx)
                cell.text = h
                cell.paragraphs[0].runs[0].bold = True

            for row_idx, item in enumerate(items, start=1):
                chk_table.cell(row_idx, 0).text = str(item.get("code", ""))
                chk_table.cell(row_idx, 1).text = str(item.get("severity", ""))
                res_str = "PASSED" if item.get("passed") else "FAILED"
                chk_table.cell(row_idx, 2).text = res_str
                chk_table.cell(row_idx, 3).text = str(item.get("message", ""))
        else:
            doc.add_paragraph("No checklist results recorded for this round.")

        doc.add_paragraph()

        # Comments Section
        doc.add_heading("Moderator Comments & Change Requests", level=2)
        comments = getattr(review, "comments", []) or []
        if comments:
            comm_table = doc.add_table(rows=len(comments) + 1, cols=4)
            comm_table.style = "Table Grid"
            c_headers = ["Category", "Target", "Comment", "Status"]
            for c_idx, ch in enumerate(c_headers):
                cell = comm_table.cell(0, c_idx)
                cell.text = ch
                cell.paragraphs[0].runs[0].bold = True

            for r_idx, c in enumerate(comments, start=1):
                cat = getattr(c.category, "value", str(c.category))
                target = getattr(c.target_type, "value", str(c.target_type))
                resolved = "RESOLVED" if getattr(c, "resolved", False) else "PENDING"
                comm_table.cell(r_idx, 0).text = cat
                comm_table.cell(r_idx, 1).text = target
                comm_table.cell(r_idx, 2).text = getattr(c, "comment", "")
                comm_table.cell(r_idx, 3).text = resolved
        else:
            doc.add_paragraph("No change requests or comments logged.")

        doc.add_paragraph()

        # Decision & Signatures
        doc.add_heading("Moderator Decision & Sign-off", level=2)
        note = getattr(review, "decision_note", None) or "None recorded."
        doc.add_paragraph(f"Decision Note: {note}")

        dec_date = getattr(review, "decided_at", None)
        date_str = dec_date.strftime("%Y-%m-%d %H:%M UTC") if dec_date else "Pending decision"
        doc.add_paragraph(f"Moderator: {moderator_name}")
        doc.add_paragraph(f"Date: {date_str}")

        out = io.BytesIO()
        doc.save(out)
        return out.getvalue()

    @classmethod
    def generate_exam_paper(
        cls,
        assessment: Any,
        seal: Optional[Any] = None,
        verification_url_base: str = "https://mindexa.edu/verify-seal",
    ) -> bytes:
        """
        Generate DOCX Exam Paper with instructions, sections, choice notes, and seal block with QR code.
        """
        doc = Document()

        # Header
        p_title = doc.add_heading(level=0)
        p_title.add_run(getattr(assessment, "title", "EXAMINATION PAPER")).bold = True
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        course_code = getattr(assessment, "course_code", "") or ""
        course_name = getattr(assessment, "course_name", "") or ""
        if course_code or course_name:
            p_course = doc.add_paragraph()
            p_course.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_course.add_run(f"Course: {course_code} - {course_name}").bold = True

        duration = getattr(assessment, "duration_minutes", None)
        total_marks = getattr(assessment, "total_marks", 100)

        p_info = doc.add_paragraph()
        p_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
        info_parts = [f"Total Marks: {total_marks}"]
        if duration:
            info_parts.append(f"Duration: {duration} minutes")
        p_info.add_run(" | ".join(info_parts))

        doc.add_paragraph()

        # Instructions
        instructions = getattr(assessment, "instructions", None)
        if instructions:
            doc.add_heading("INSTRUCTIONS TO CANDIDATES", level=2)
            doc.add_paragraph(instructions)
            doc.add_paragraph()

        # Sections and Questions
        all_secs = getattr(assessment, "sections", []) or []
        secs = sorted(
            [s for s in all_secs if not getattr(s, "is_deleted", False)],
            key=lambda x: getattr(x, "order_index", 0),
        )

        all_aqs = getattr(assessment, "assessment_questions", []) or []
        aqs = [aq for aq in all_aqs if not getattr(aq, "is_deleted", False)]

        question_counter = 1

        if secs:
            for sec in secs:
                doc.add_heading(sec.title, level=1)
                # Choice note
                n_choice = getattr(sec, "questions_to_answer", None)
                sec_aqs = [
                    aq for aq in aqs if getattr(aq, "assessment_section_id", None) == sec.id
                ]
                sec_aqs.sort(key=lambda x: getattr(x, "order_index", 0))

                if n_choice is not None and len(sec_aqs) > n_choice:
                    doc.add_paragraph(
                        f"[NOTE: Answer any {n_choice} of the {len(sec_aqs)} questions in this section.]",
                    ).italic = True
                elif n_choice is not None:
                    doc.add_paragraph(f"[NOTE: Answer {n_choice} questions.]").italic = True
                else:
                    doc.add_paragraph("[NOTE: All questions in this section are compulsory.]").italic = True

                for aq in sec_aqs:
                    q = getattr(aq, "question", None)
                    marks = _effective_aq_marks(aq)
                    q_text = (q.content if q else "") or "Question stem"
                    p_q = doc.add_paragraph()
                    p_q.add_run(f"Question {question_counter} ").bold = True
                    p_q.add_run(f"({marks} marks)\n")
                    p_q.add_run(q_text)

                    # Options if MCQ
                    if q and getattr(q, "options", None):
                        q_type = getattr(q, "question_type", None)
                        if hasattr(q_type, "value"):
                            q_type = q_type.value
                        if q_type in (QuestionType.MCQ.value, QuestionType.TRUE_FALSE.value):
                            opt_list = sorted(
                                [o for o in q.options if not getattr(o, "is_deleted", False)],
                                key=lambda x: getattr(x, "order_index", 0),
                            )
                            for opt_idx, opt in enumerate(opt_list):
                                label = chr(65 + opt_idx)  # A, B, C, D
                                txt = opt.content or opt.option_text or ""
                                doc.add_paragraph(f"    ({label}) {txt}")

                    doc.add_paragraph()
                    question_counter += 1
        else:
            # Unsectioned questions
            for aq in aqs:
                q = getattr(aq, "question", None)
                marks = _effective_aq_marks(aq)
                q_text = (q.content if q else "") or "Question stem"
                p_q = doc.add_paragraph()
                p_q.add_run(f"Question {question_counter} ").bold = True
                p_q.add_run(f"({marks} marks)\n")
                p_q.add_run(q_text)
                doc.add_paragraph()
                question_counter += 1

        # Seal block at the end if sealed
        if seal:
            doc.add_page_break()
            doc.add_heading("EXAM OFFICE OFFICIAL SEAL & VERIFICATION", level=2)
            code = getattr(seal, "verification_code", "")
            sealed_at = getattr(seal, "sealed_at", datetime.now(UTC))
            date_str = sealed_at.strftime("%Y-%m-%d %H:%M UTC")

            p_seal = doc.add_paragraph()
            p_seal.add_run(f"Official Verification Code: {code}\n").bold = True
            p_seal.add_run(f"Sealed At: {date_str}\n")
            p_seal.add_run(f"Content Hash (SHA-256): {getattr(seal, 'content_hash', '')[:16]}...\n")

            # Generate QR code
            verify_url = f"{verification_url_base.rstrip('/')}/{code}"
            if qrcode is not None:
                qr = qrcode.QRCode(box_size=4, border=2)
                qr.add_data(verify_url)
                qr.make(fit=True)
                img = qr.make_image(fill_color="black", back_color="white")
                qr_buf = io.BytesIO()
                img.save(qr_buf, format="PNG")
                qr_buf.seek(0)

                doc.add_picture(qr_buf, width=Inches(1.8))
                p_qr_caption = doc.add_paragraph("Scan QR code or visit Mindexa portal to verify paper authenticity.")
                p_qr_caption.italic = True
            else:
                p_qr_caption = doc.add_paragraph(f"Verify paper authenticity at: {verify_url}")
                p_qr_caption.italic = True

        out = io.BytesIO()
        doc.save(out)
        return out.getvalue()

    @classmethod
    def generate_marking_scheme(
        cls,
        assessment: Any,
    ) -> bytes:
        """
        Generate DOCX Marking Scheme Document.
        Includes: assessment metadata, question stems, allocated marks, and answers/explanations.
        """
        doc = Document()

        p_title = doc.add_heading(level=0)
        p_title.add_run(f"OFFICIAL MARKING SCHEME & GUIDE\n{getattr(assessment, 'title', '')}").bold = True
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        course_code = getattr(assessment, "course_code", "") or ""
        course_name = getattr(assessment, "course_name", "") or ""
        if course_code or course_name:
            p_course = doc.add_paragraph()
            p_course.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_course.add_run(f"Course: {course_code} - {course_name}").bold = True

        doc.add_paragraph(f"Total Marks: {getattr(assessment, 'total_marks', 100)}")
        doc.add_paragraph()

        all_secs = getattr(assessment, "sections", []) or []
        secs = sorted(
            [s for s in all_secs if not getattr(s, "is_deleted", False)],
            key=lambda x: getattr(x, "order_index", 0),
        )

        all_aqs = getattr(assessment, "assessment_questions", []) or []
        aqs = [aq for aq in all_aqs if not getattr(aq, "is_deleted", False)]

        q_counter = 1

        if secs:
            for sec in secs:
                doc.add_heading(sec.title, level=1)
                sec_aqs = [
                    aq for aq in aqs if getattr(aq, "assessment_section_id", None) == sec.id
                ]
                sec_aqs.sort(key=lambda x: getattr(x, "order_index", 0))

                for aq in sec_aqs:
                    q = getattr(aq, "question", None)
                    marks = _effective_aq_marks(aq)
                    q_text = (q.content if q else "") or "Question stem"
                    p_q = doc.add_paragraph()
                    p_q.add_run(f"Question {q_counter} ").bold = True
                    p_q.add_run(f"[{marks} marks]\n")
                    p_q.add_run(f"Stem: {q_text}\n")

                    # Correct answer / marking guide
                    p_ans = doc.add_paragraph()
                    p_ans.add_run("Marking Criteria / Expected Answer:\n").bold = True

                    if q and getattr(q, "options", None):
                        correct_opts = [
                            o for o in q.options
                            if not getattr(o, "is_deleted", False) and getattr(o, "is_correct", False)
                        ]
                        if correct_opts:
                            p_ans.add_run("Correct Answer: " + ", ".join(
                                o.content or o.option_text or "" for o in correct_opts
                            ) + "\n")

                    explanation = getattr(q, "explanation", None) if q else None
                    if explanation:
                        p_ans.add_run(f"Explanation/Rubric Guide: {explanation}\n")
                    else:
                        p_ans.add_run("Standard evaluation based on question objectives and method accuracy.\n")

                    doc.add_paragraph()
                    q_counter += 1
        else:
            for aq in aqs:
                q = getattr(aq, "question", None)
                marks = _effective_aq_marks(aq)
                q_text = (q.content if q else "") or "Question stem"
                p_q = doc.add_paragraph()
                p_q.add_run(f"Question {q_counter} ").bold = True
                p_q.add_run(f"[{marks} marks]\n")
                p_q.add_run(f"Stem: {q_text}\n")
                explanation = getattr(q, "explanation", None) if q else None
                if explanation:
                    doc.add_paragraph(f"Marking Guide: {explanation}")
                else:
                    doc.add_paragraph("Marking Guide: Standard marking guide.")
                doc.add_paragraph()
                q_counter += 1

        out = io.BytesIO()
        doc.save(out)
        return out.getvalue()
