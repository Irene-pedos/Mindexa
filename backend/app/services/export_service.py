"""
app/services/export_service.py

Phase 6C: Excel export engine using openpyxl.

Two export types:
  1. ASSESSMENT_MARKSHEET  — per-assessment mark sheet with Summary + Meta sheets
  2. COURSE_CONSOLIDATED   — per-course mark sheet using Phase 6A unweighted average

Shared rules:
  - sanitize_cell() on all user-sourced strings (injection protection)
  - Numbers as numeric cells; percentages rounded 2dp
  - Freeze header row, column widths, landscape print setup
  - ≤300 rows → synchronous result; >300 rows → Celery ExportJob + polling
  - Every export [AUDIT]
  - NO integrity flags, NO AI details exported
"""
from __future__ import annotations

import io
import uuid
from typing import Any

import structlog
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.exceptions import NotFoundError, ValidationError
from app.db.base import utcnow
from app.db.enums import ExportJobStatus, ExportType
from app.db.models.auth import User
from app.utils.spreadsheet import sanitize_cell

log = structlog.get_logger(__name__)

_SYNC_ROW_LIMIT = 300
_EXPORT_TTL_HOURS = 24


class ExportService:
    def __init__(self, db: AsyncSession) -> None:
        self.db = db

    # ── Assessment Marksheet ──────────────────────────────────────────────────

    async def export_assessment_marksheet(
        self,
        assessment_id: uuid.UUID,
        requester: User,
        template_code: str = "default_marksheet",
    ) -> bytes:
        """
        Build the Assessment Mark Sheet Excel workbook and return raw bytes.
        For >300 rows, call create_async_export_job() instead.

        Sheets: "Marks", "Summary", "Meta"
        Rule: NO integrity flags, NO AI details.
        """
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise ValidationError(
                message="openpyxl is required for Excel exports. Install it with: pip install openpyxl",
                field="export",
            )

        from app.db.models.assessment import Assessment
        from app.db.models.auth import User as UserModel
        from app.db.models.question import AssessmentQuestion
        from app.db.models.result import AssessmentResult

        # Load assessment
        assessment_res = await self.db.execute(
            select(Assessment)
            .where(Assessment.id == assessment_id, Assessment.is_deleted == False)  # noqa: E712
            .options(
                selectinload(Assessment.sections),
                selectinload(Assessment.assessment_questions).selectinload(AssessmentQuestion.question),
            )
        )
        assessment = assessment_res.scalars().first()
        if not assessment:
            raise NotFoundError(resource="Assessment", resource_id=str(assessment_id))

        # Load released results with attempt and student info
        results_res = await self.db.execute(
            select(AssessmentResult)
            .where(
                AssessmentResult.assessment_id == assessment_id,
                AssessmentResult.is_released == True,  # noqa: E712
                AssessmentResult.is_deleted == False,  # noqa: E712
            )
            .options(
                selectinload(AssessmentResult.attempt),
            )
            .order_by(AssessmentResult.student_id)
        )
        results = list(results_res.scalars().all())

        # Load student profiles for display
        student_ids = [r.student_id for r in results]
        students: dict[uuid.UUID, UserModel] = {}
        if student_ids:
            students_res = await self.db.execute(
                select(UserModel)
                .where(UserModel.id.in_(student_ids))
                .options(selectinload(UserModel.profile))
            )
            for s in students_res.scalars().all():
                students[s.id] = s

        wb = openpyxl.Workbook()

        # ── Sheet 1: Marks ────────────────────────────────────────────────────
        ws_marks = wb.active
        ws_marks.title = "Marks"

        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_font = Font(color="FFFFFF", bold=True)

        from app.services.assessment_marks import required_marks
        req_marks = required_marks(assessment) or getattr(assessment, "total_marks", 0)

        headers = [
            "No.", "Reg No", "Full Name",
            "Raw Total", f"Max ({req_marks})",
            "Percentage", "Grade", "Status", "Marker",
        ]
        for col_idx, h in enumerate(headers, 1):
            cell = ws_marks.cell(row=1, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        ws_marks.freeze_panes = "A2"

        row_num = 2
        for i, result in enumerate(results, 1):
            student = students.get(result.student_id)
            profile = student.profile if student and hasattr(student, "profile") else None
            full_name = sanitize_cell(
                f"{profile.first_name} {profile.last_name}" if profile else str(result.student_id)
            )
            reg_no = sanitize_cell(
                getattr(profile, "registration_number", "") if profile else ""
            )
            grader_name = ""
            grader_id = getattr(result.attempt, "graded_by_id", None) if result.attempt else None
            if grader_id:
                grader = students.get(grader_id)
                if grader and hasattr(grader, "profile") and grader.profile:
                    grader_name = sanitize_cell(
                        f"{grader.profile.first_name} {grader.profile.last_name}"
                    )

            ws_marks.append([
                i,
                reg_no,
                full_name,
                round(result.total_score, 2),
                round(result.max_score, 2),
                round(result.percentage, 2),
                sanitize_cell(result.letter_grade.value if result.letter_grade else ""),
                "PASS" if result.is_passing else "FAIL",
                grader_name,
            ])
            row_num += 1

        # Auto column widths
        for col in ws_marks.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws_marks.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 40)

        ws_marks.page_setup.orientation = "landscape"

        # ── Sheet 2: Summary ──────────────────────────────────────────────────
        ws_summary = wb.create_sheet("Summary")
        percentages = [r.percentage for r in results if r.percentage is not None]
        scores = [r.total_score for r in results]
        passed_count = sum(1 for r in results if r.is_passing)

        _summary_rows = [
            ("Total Students", len(results)),
            ("Passed", passed_count),
            ("Failed", len(results) - passed_count),
            ("Pass Rate (%)", round(passed_count / len(results) * 100, 2) if results else 0),
            ("Average (%)", round(sum(percentages) / len(percentages), 2) if percentages else 0),
            ("Highest (%)", round(max(percentages), 2) if percentages else 0),
            ("Lowest (%)", round(min(percentages), 2) if percentages else 0),
        ]
        for label, value in _summary_rows:
            ws_summary.append([label, value])

        # ── Sheet 3: Meta ─────────────────────────────────────────────────────
        ws_meta = wb.create_sheet("Meta")
        import hashlib

        export_ts = utcnow().isoformat()
        seal_input = f"{assessment_id}:{len(results)}:{export_ts}"
        seal_hash = hashlib.sha256(seal_input.encode()).hexdigest()[:16].upper()

        ws_meta.append(["Assessment ID", str(assessment_id)])
        ws_meta.append(["Title", sanitize_cell(assessment.title)])
        ws_meta.append(["Total Marks", req_marks])
        ws_meta.append(["Exported At (UTC)", export_ts])
        ws_meta.append(["Verification Code", seal_hash])
        ws_meta.append(["Row Count", len(results)])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        log.info(
            "[AUDIT] export.assessment_marksheet",
            assessment_id=str(assessment_id),
            requester_id=str(requester.id),
            row_count=len(results),
        )
        return buf.read()

    # ── Course Consolidated Marksheet ─────────────────────────────────────────

    async def export_course_consolidated(
        self,
        course_id: uuid.UUID,
        period_id: uuid.UUID,
        requester: User,
        template_code: str = "default_course_marksheet",
    ) -> bytes:
        """
        Course Consolidated Mark Sheet: one row per student.
        Uses Phase 6A unweighted coursework average (R-20) as the single source of truth.

        Columns: Reg No, Full Name, each CAT %, each formative %,
                 Coursework Average, Summative %, Eligibility, Course Outcome, Retake
        """
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise ValidationError(
                message="openpyxl is required for Excel exports.",
                field="export",
            )

        from app.db.enums import AssessmentType
        from app.db.models.academic import Course, StudentEnrollment
        from app.db.models.assessment import Assessment
        from app.db.models.auth import User as UserModel
        from app.db.models.result import AssessmentResult
        from app.services.eligibility_rules import CourseworkMarkItem, coursework_average

        # Load course
        course_res = await self.db.execute(
            select(Course).where(Course.id == course_id, Course.is_deleted == False)  # noqa: E712
        )
        course = course_res.scalars().first()
        if not course:
            raise NotFoundError(resource="Course", resource_id=str(course_id))

        # Load enrolled students
        enrollments_res = await self.db.execute(
            select(StudentEnrollment)
            .where(
                StudentEnrollment.course_id == course_id,
                StudentEnrollment.is_deleted == False,  # noqa: E712
            )
            .options(selectinload(StudentEnrollment.student))
        )
        enrollments = list(enrollments_res.scalars().all())

        # Load assessments for this course/period
        assessments_res = await self.db.execute(
            select(Assessment).where(
                Assessment.course_id == course_id,
                Assessment.academic_period_id == period_id,
                Assessment.is_deleted == False,  # noqa: E712
            ).order_by(Assessment.created_at)
        )
        assessments = list(assessments_res.scalars().all())

        cats = [a for a in assessments if getattr(a, "assessment_type", None) in ("CAT", AssessmentType.CAT)]
        formatives = [a for a in assessments if getattr(a, "assessment_type", None) in ("FORMATIVE", AssessmentType.FORMATIVE)]
        summatives = [a for a in assessments if getattr(a, "assessment_type", None) in ("SUMMATIVE", AssessmentType.SUMMATIVE)]

        # Load all results for this course/period indexed by (student_id, assessment_id)
        all_assessment_ids = [a.id for a in assessments]
        results_map: dict[tuple, AssessmentResult] = {}
        if all_assessment_ids:
            results_res = await self.db.execute(
                select(AssessmentResult).where(
                    AssessmentResult.assessment_id.in_(all_assessment_ids),
                    AssessmentResult.is_released == True,  # noqa: E712
                    AssessmentResult.is_deleted == False,  # noqa: E712
                )
            )
            for r in results_res.scalars().all():
                results_map[(r.student_id, r.assessment_id)] = r

        # Load student profiles
        student_ids = [e.student_id for e in enrollments]
        students: dict[uuid.UUID, UserModel] = {}
        if student_ids:
            students_res = await self.db.execute(
                select(UserModel)
                .where(UserModel.id.in_(student_ids))
                .options(selectinload(UserModel.profile))
            )
            for s in students_res.scalars().all():
                students[s.id] = s

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Consolidated"

        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_font = Font(color="FFFFFF", bold=True)

        # Build dynamic headers
        headers = ["Reg No", "Full Name"]
        for i, cat in enumerate(cats, 1):
            headers.append(sanitize_cell(f"CAT{i} (%)"))
        for i, form in enumerate(formatives, 1):
            headers.append(sanitize_cell(f"Form{i} (%)"))
        headers += ["CW Average (%)", "Summative (%)", "Eligibility", "Outcome", "Retake?"]

        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        ws.freeze_panes = "A2"

        for enrollment in enrollments:
            student = students.get(enrollment.student_id)
            profile = student.profile if student and hasattr(student, "profile") else None
            full_name = sanitize_cell(
                f"{profile.first_name} {profile.last_name}" if profile else str(enrollment.student_id)
            )
            reg_no = sanitize_cell(
                getattr(profile, "registration_number", "") if profile else ""
            )

            row: list[Any] = [reg_no, full_name]

            cat_pcts: list[CourseworkMarkItem] = []
            for cat in cats:
                result = results_map.get((enrollment.student_id, cat.id))
                pct = round(result.percentage, 2) if result else None
                row.append(pct if pct is not None else "")
                cat_pcts.append(CourseworkMarkItem(
                    percentage=result.percentage if result else None,
                    is_ungraded=result is None,
                ))

            form_pcts: list[CourseworkMarkItem] = []
            for form in formatives:
                result = results_map.get((enrollment.student_id, form.id))
                pct = round(result.percentage, 2) if result else None
                row.append(pct if pct is not None else "")
                form_pcts.append(CourseworkMarkItem(
                    percentage=result.percentage if result else None,
                    is_ungraded=result is None,
                ))

            # Coursework average via Phase 6A function (R-20 unweighted)
            cw_avg = coursework_average(cat_pcts + form_pcts)
            row.append(round(cw_avg, 2) if cw_avg is not None else "")

            # Summative
            summative_pct = None
            if summatives:
                summative_result = results_map.get((enrollment.student_id, summatives[0].id))
                summative_pct = round(summative_result.percentage, 2) if summative_result else None
            row.append(summative_pct if summative_pct is not None else "")

            # Eligibility status
            row.append("ELIGIBLE" if cw_avg is not None and cw_avg >= 50.0 else "BLOCKED")

            # Course outcome
            if summative_pct is not None and cw_avg is not None:
                outcome = "PASS" if summative_pct >= 50.0 else "FAIL"
            else:
                outcome = "PENDING"
            row.append(outcome)

            # Retake flag
            row.append("Yes" if outcome == "FAIL" else "No")

            ws.append(row)

        # Column widths + landscape
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 40)
        ws.page_setup.orientation = "landscape"

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        log.info(
            "[AUDIT] export.course_consolidated",
            course_id=str(course_id),
            period_id=str(period_id),
            requester_id=str(requester.id),
            row_count=len(enrollments),
        )
        return buf.read()

    # ── Async Job (>300 rows) ─────────────────────────────────────────────────

    async def create_async_export_job(
        self,
        requester: User,
        export_type: ExportType,
        assessment_id: uuid.UUID | None = None,
        course_id: uuid.UUID | None = None,
        period_id: uuid.UUID | None = None,
        template_code: str = "default_marksheet",
    ) -> Any:
        """Create an ExportJob and enqueue it to Celery. Returns the job record."""
        from app.db.repositories.export_repo import ExportRepository

        repo = ExportRepository(self.db)
        template = await repo.get_template_by_code(template_code)
        if not template:
            raise NotFoundError(resource="ExportTemplate", resource_id=template_code)

        job = await repo.create_job(
            requested_by_id=requester.id,
            export_type=export_type,
            template_id=template.id,
            assessment_id=assessment_id,
            course_id=course_id,
            period_id=period_id,
        )

        # Enqueue Celery task
        try:
            from app.workers.tasks.export_tasks import run_export_job
            run_export_job.apply_async(args=[str(job.id)], queue="exports")
        except Exception as exc:
            log.warning("export.celery_enqueue_failed", job_id=str(job.id), error=str(exc))
            # Mark job FAILED if we can't enqueue
            await repo.update_job_status(job, ExportJobStatus.FAILED, error=str(exc))

        log.info(
            "[AUDIT] export.job_created",
            job_id=str(job.id),
            export_type=export_type.value,
            requester_id=str(requester.id),
        )
        return job
